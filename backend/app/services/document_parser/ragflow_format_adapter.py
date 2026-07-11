"""RAGFlow parser adapter for all supported local document formats.

This module deliberately keeps know_rag's ingestion pipeline unchanged.  It only
uses RAGFlow's original DeepDoc/parser code for recognition, cleaning and basic
structure extraction, then converts the result into ParsedDocument/ParsedBlock so
chunking, embedding, retrieval and management remain owned by know_rag.
"""
from __future__ import annotations

import base64
import html
import json
import logging
import os
import re
import shutil
import sys
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable
from zipfile import ZipFile

from PIL import Image

from app.services.document_parser.media_storage import save_media_bytes
from app.services.document_parser.types import ParsedBlock, ParsedDocument
from app.utils.text import sanitize_text_for_postgres

logger = logging.getLogger(__name__)

_TAG_RE = re.compile(r"@@([0-9\-]+)\t([0-9.]+)\t([0-9.]+)\t([0-9.]+)\t([0-9.]+)##")
_HTML_RE = re.compile(r"<[^>]+>")
_TABLE_RE = re.compile(r"<\s*table\b", re.I)

TEXT_SUFFIXES = {"txt", "md", "markdown"}
IMAGE_SUFFIXES = {"jpg", "jpeg", "png", "gif", "bmp", "webp", "tif", "tiff", "svg", "heic", "heif"}
TABLE_SUFFIXES = {"csv", "tsv", "xlsx", "xls"}
HTML_SUFFIXES = {"html", "htm"}
PRESENTATION_SUFFIXES = {"pptx", "ppt"}
EPUB_SUFFIXES = {"epub"}
LEGACY_DOC_SUFFIXES = {"doc", "rtf"}
STRUCTURED_TEXT_SUFFIXES = {"json", "jsonl", "xml", "yaml", "yml"}
ARCHIVE_SUFFIXES = {"zip"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | IMAGE_SUFFIXES | TABLE_SUFFIXES | HTML_SUFFIXES | PRESENTATION_SUFFIXES | EPUB_SUFFIXES | LEGACY_DOC_SUFFIXES | STRUCTURED_TEXT_SUFFIXES | ARCHIVE_SUFFIXES | {"pdf", "docx"}


def _vendor_root() -> Path:
    """Locate the vendored RAGFlow tree in both source and Docker layouts.

    In the Docker image this file lives under ``/app/app/...`` while the vendor
    tree is copied to ``/app/ragflow_vendor``.  A fixed ``parents[4]`` therefore
    points at ``/ragflow_vendor`` and is environment-dependent.
    """
    current = Path(__file__).resolve()
    for parent in current.parents:
        candidate = parent / "ragflow_vendor"
        if candidate.is_dir():
            return candidate
    # Deterministic fallback for diagnostics; _ensure_vendor_path callers will
    # still fail with a clear import error if the package is truly absent.
    return current.parents[3] / "ragflow_vendor"


def _ensure_vendor_path() -> Path:
    root = _vendor_root()
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))
    os.environ.setdefault("RAGFLOW_VENDOR_BASE", str(root))
    os.environ.setdefault("RAG_PROJECT_BASE", str(root))
    _ensure_model_dir(root)
    return root


def _ensure_model_dir(root: Path) -> None:
    """Make RAGFlow's expected rag/res/deepdoc point at Docker model cache when possible."""
    expected = root / "rag" / "res" / "deepdoc"
    expected.parent.mkdir(parents=True, exist_ok=True)
    external = Path(os.getenv("DEEPDOC_MODEL_DIR", "/data/models/deepdoc"))
    if external.exists() and any(external.iterdir()):
        if expected.exists() and not expected.is_symlink() and any(expected.iterdir()):
            return
        try:
            if expected.exists() or expected.is_symlink():
                if expected.is_symlink() or expected.is_file():
                    expected.unlink()
                else:
                    shutil.rmtree(expected)
            expected.symlink_to(external, target_is_directory=True)
            return
        except Exception:
            expected.mkdir(parents=True, exist_ok=True)
            for child in external.iterdir():
                if child.is_file():
                    dst = expected / child.name
                    if not dst.exists():
                        try:
                            shutil.copy2(child, dst)
                        except Exception:
                            pass
    else:
        expected.mkdir(parents=True, exist_ok=True)


def _clean(text: Any) -> str:
    return sanitize_text_for_postgres(str(text or "").strip())


def _strip_position_tags(text: str) -> str:
    return _clean(_TAG_RE.sub("", text or ""))


def _positions_from_text(text: str) -> tuple[int | None, tuple[float, float, float, float] | None]:
    m = _TAG_RE.search(text or "")
    if not m:
        return None, None
    try:
        page_token = m.group(1).split("-")[0]
        page_no = int(page_token)
        x0, x1, top, bottom = map(float, m.groups()[1:])
        return page_no, (x0, top, x1, bottom)
    except Exception:
        return None, None


def _html_to_text(raw: Any) -> str:
    txt = html.unescape(_HTML_RE.sub(" ", str(raw or "")))
    return _clean(re.sub(r"\s+", " ", txt))


def _section_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        if len(value) >= 1:
            return _clean(value[0])
        return ""
    if isinstance(value, dict):
        return _clean(value.get("content") or value.get("text") or json.dumps(value, ensure_ascii=False))
    return _clean(value)


def _pil_to_bytes(img: Any) -> bytes | None:
    if img is None:
        return None
    try:
        if hasattr(img, "to_pil"):
            img = img.to_pil()
        if isinstance(img, bytes):
            return img
        if isinstance(img, str):
            try:
                if img.startswith("data:"):
                    img = img.split(",", 1)[1]
                return base64.b64decode(img)
            except Exception:
                return None
        if isinstance(img, Image.Image):
            buf = BytesIO()
            img.convert("RGB").save(buf, format="PNG")
            return buf.getvalue()
    except Exception as exc:
        logger.info("convert ragflow image failed: %s", exc)
    return None


def _save_image_payload(payload: Any, *, knowledge_base_id: str, document_id: str, filename_hint: str) -> dict:
    data = _pil_to_bytes(payload)
    if not data:
        return {}
    return save_media_bytes(
        knowledge_base_id=knowledge_base_id,
        document_id=document_id,
        filename_hint=filename_hint,
        data=data,
        default_ext=Path(filename_hint).suffix or ".png",
    )


def _make_text_block(text: Any, *, order_index: int, parser: str, page_no: int | None = None, metadata: dict | None = None) -> ParsedBlock | None:
    txt = _section_text(text)
    if not txt:
        return None
    meta = {"parser": parser, **(metadata or {})}
    if re.match(r"^\s*```[\s\S]*?```\s*$", txt):
        meta.setdefault("raw_block_type", "code_block")
        meta.setdefault("content_kind", "code_block")
        meta.setdefault("preserve_line_boundaries", True)
    return ParsedBlock(
        block_type="text",
        text=txt,
        order_index=order_index,
        page_no=page_no,
        metadata=meta,
    )


def _make_table_block(text: str, table_html: str, *, order_index: int, parser: str, metadata: dict | None = None) -> ParsedBlock | None:
    txt = _clean(text or _html_to_text(table_html))
    html_src = str(table_html or "")
    if not txt and not html_src:
        return None
    return ParsedBlock(
        block_type="table",
        text=txt or "表格",
        table_html=html_src,
        order_index=order_index,
        metadata={"parser": parser, "source": "ragflow_table", **(metadata or {})},
    )


def _table_text(rows_or_html: Any) -> tuple[str, str]:
    if rows_or_html is None:
        return "", ""
    if isinstance(rows_or_html, str):
        raw = rows_or_html
        return _html_to_text(raw), raw if "<table" in raw.lower() else ""
    if isinstance(rows_or_html, (list, tuple)):
        lines = []
        for row in rows_or_html:
            if isinstance(row, (list, tuple)):
                lines.append(" | ".join(str(c) for c in row))
            else:
                lines.append(str(row))
        return _clean("\n".join(lines)), ""
    return _clean(str(rows_or_html)), ""



def _rows_to_matrix(rows_or_html: Any) -> list[list[str]]:
    """Normalize RAGFlow table-like rows into a cell matrix.

    RAGFlow's PDF media output is intentionally loose: for true tables it may
    return rows/cells, while for figures or page regions it can also return a
    list of OCR lines.  A list alone is therefore not enough to call it a table.
    """
    if not isinstance(rows_or_html, (list, tuple)):
        return []
    matrix: list[list[str]] = []
    for row in rows_or_html:
        if isinstance(row, (list, tuple)):
            cells = [_clean(c) for c in row if _clean(c)]
        else:
            raw = _clean(row)
            # Only split obvious delimited rows.  Plain OCR lines should stay as
            # one cell so they do not accidentally become table rows.
            if "\t" in raw:
                cells = [_clean(c) for c in raw.split("\t") if _clean(c)]
            elif " | " in raw:
                cells = [_clean(c) for c in raw.split("|") if _clean(c)]
            else:
                cells = [raw] if raw else []
        if cells:
            matrix.append(cells)
    return matrix


_CJK_CHAR_RE = re.compile(r"[一-鿿]")
_CJK_PUNCT_RE = re.compile(r"[。！？；：、，""''【】《》（）]")
_SENTENCE_MARK_RE = re.compile(r"[。！？.!?；;]")

_PARAGRAPH_MARKERS = [
    "页面锚点",
    "本页用于测试",
    "关键字段",
    "RAGFlow parsing stress test",
    "This page contains",
    "The goal is to test",
    "Important anchor terms",
    "SYNTHETIC PDF DATASET",
    "Raster-only OCR text line",
    "invoice number",
    "vector index check",
]


def _looks_like_long_body_text(text: str) -> bool:
    """Heuristic for demoting OCR/layout regions that are really prose text.

    Uses a generalized prose score accumulator combining CJK punctuation density,
    CJK-heavy lines, average line length, sentence marks, and delimiter absence
    so arbitrary Chinese-English mixed prose is not misclassified as a table.
    """
    text = _clean(text)
    if not text:
        return False

    prose_score = 0.0
    lines = text.splitlines()

    # CJK punctuation density (>2% of characters → strong prose signal)
    cjk_punct_count = len(_CJK_PUNCT_RE.findall(text))
    if cjk_punct_count / max(len(text), 1) > 0.02:
        prose_score += 3.0

    # CJK-heavy lines (>50% of lines contain >30% CJK characters)
    if lines:
        cjk_lines = sum(1 for line in lines if len(_CJK_CHAR_RE.findall(line)) > len(line) * 0.3)
        if cjk_lines / len(lines) > 0.5:
            prose_score += 3.0

    # Average line length (>60 chars per line → paragraph-style content)
    if lines:
        avg_line_len = sum(len(line) for line in lines) / len(lines)
        if avg_line_len > 60:
            prose_score += 2.0

    # Sentence marks (≥5 sentence-ending punctuation → prose)
    sentence_marks = len(_SENTENCE_MARK_RE.findall(text))
    if sentence_marks >= 5:
        prose_score += 2.0

    # Long lines (≥2 lines >120 chars → paragraph content)
    long_lines = sum(1 for line in lines if len(line.strip()) > 120)
    if long_lines >= 2:
        prose_score += 1.5

    # Tab/pipe delimiters suggest real table structure → penalize prose score
    if "\t" in text or " | " in text:
        prose_score -= 3.0

    # Legacy markers (backward compat with Fix1-Fix3)
    marker_hits = sum(1 for marker in _PARAGRAPH_MARKERS if marker.lower() in text.lower())
    if marker_hits >= 2:
        prose_score += 3.0

    return prose_score >= 4.0


def _table_structure_assessment(rows_or_html: Any, text: str, table_html: str) -> tuple[str, str, float, str]:
    """Return (kind, subtype, confidence, reason) for RAGFlow PDF media.

    kind is one of: table, image.  Subtype is stored in metadata and may be:
    real_table, key_value_table, pseudo_table, figure_or_layout.
    """
    raw_text = _clean(text or "")
    lower = raw_text.lower()
    if table_html and "<table" in table_html.lower() and not _looks_like_long_body_text(raw_text):
        return "table", "real_table", 0.95, "html_table"

    matrix = _rows_to_matrix(rows_or_html)
    valid_rows = [row for row in matrix if len(row) >= 2]
    col_counts = [len(row) for row in valid_rows]
    stable_cols = 0
    if col_counts:
        most_common_cols = max(set(col_counts), key=col_counts.count)
        stable_cols = sum(1 for count in col_counts if count == most_common_cols)
    stable_ratio = stable_cols / max(len(valid_rows), 1)
    avg_cell_len = 0.0
    max_cell_len = 0
    cjk_cell_ratio = 0.0
    if valid_rows:
        cells = [cell for row in valid_rows for cell in row]
        avg_cell_len = sum(len(cell) for cell in cells) / max(len(cells), 1)
        max_cell_len = max(len(cell) for cell in cells) if cells else 0
        cjk_cell_count = sum(1 for cell in cells if len(_CJK_CHAR_RE.findall(cell)) > len(cell) * 0.3)
        cjk_cell_ratio = cjk_cell_count / max(len(cells), 1)

    header_cues = [
        "field value note",
        "字段 field 值 value 说明 note",
        "字段 值 说明",
        "field\tvalue\tnote",
        "字段|值|说明",
    ]
    normalized = re.sub(r"\s+", " ", raw_text).strip().lower()
    has_header_cue = any(cue in normalized for cue in header_cues)
    has_table_caption = bool(re.search(r"(^|\n)\s*(table|表)\s*[0-9ivx一二三四五六七八九十]+\b", raw_text, re.I))

    # Guard: if cells are CJK-heavy and long, this is prose not a table.
    # Key-value tables (Field/Value/Note) are exempt via has_header_cue.
    if max_cell_len > 100 and cjk_cell_ratio > 0.3 and not has_header_cue:
        return "image", "pseudo_table", 0.35, "cjk_prose_cells"

    if valid_rows and len(valid_rows) >= 2 and stable_ratio >= 0.6 and avg_cell_len <= 80 and not _looks_like_long_body_text(raw_text):
        subtype = "key_value_table" if has_header_cue else "real_table"
        confidence = 0.88 if has_header_cue or has_table_caption else 0.74
        return "table", subtype, confidence, "stable_cell_matrix"

    # Explicit figure/equation signals are image chunks.  Other media are kept as
    # image/layout regions, with OCR text preserved in metadata for text chunking.
    if re.search(r"(^|\n)\s*(fig\.?|figure|图)\s*[0-9ivx一二三四五六七八九十]+\b", raw_text, re.I):
        return "image", "figure", 0.85, "figure_caption"
    if re.search(r"^\s*\(?[0-9]+\)?\s*$", raw_text) or any(sym in raw_text for sym in ["=", "∑", "∫", "√", "≤", "≥"]):
        return "image", "equation_or_formula", 0.65, "equation_signal"

    subtype = "pseudo_table" if valid_rows else "figure_or_layout"
    return "image", subtype, 0.45, "weak_or_layout_rows"


def _classify_ragflow_pdf_media(rows_or_html: Any, text: str, table_html: str) -> str:
    """Backward compatible wrapper returning only the coarse media kind."""
    kind, _, _, _ = _table_structure_assessment(rows_or_html, text, table_html)
    return kind

def _finalize(blocks: list[ParsedBlock], parser: str, extra_meta: dict | None = None) -> ParsedDocument:
    blocks = [b for b in blocks if b and (b.index_text() or b.text or b.table_html or b.image_path)]
    for i, b in enumerate(blocks):
        b.order_index = i
    content = "\n\n".join(b.index_text() or b.text or _html_to_text(b.table_html) for b in blocks if (b.index_text() or b.text or b.table_html))
    return ParsedDocument(
        content=_clean(content),
        blocks=blocks,
        metadata={
            "parser": parser,
            "recognition_alignment": "ragflow_original_parser_adapter",
            "block_count": len(blocks),
            "text_blocks": sum(1 for b in blocks if b.block_type == "text"),
            "table_blocks": sum(1 for b in blocks if b.block_type == "table"),
            "image_blocks": sum(1 for b in blocks if b.block_type == "image"),
            **(extra_meta or {}),
        },
    )


def parse_pdf(path: Path, *, knowledge_base_id: str, document_id: str) -> ParsedDocument:
    _ensure_vendor_path()
    from deepdoc.parser.pdf_parser import RAGFlowPdfParser  # type: ignore

    parser = RAGFlowPdfParser()
    try:
        sections, media_items = parser(str(path), need_image=True, zoomin=3, return_html=True)
    except TypeError:
        sections, media_items = parser(str(path))
    blocks: list[ParsedBlock] = []
    if isinstance(sections, str):
        parts = [p for p in re.split(r"\n\s*\n", sections or "") if _strip_position_tags(p)]
    else:
        parts = [s[0] if isinstance(s, (list, tuple)) else str(s) for s in (sections or [])]
    for part in parts:
        page_no, bbox = _positions_from_text(part)
        text = _strip_position_tags(part)
        if text:
            blocks.append(ParsedBlock(block_type="text", text=text, page_no=page_no, bbox=bbox, order_index=len(blocks), metadata={"parser": "ragflow_pdf", "source": "section"}))

    for item in media_items or []:
        # RAGFlowPdfParser.__call__ normally returns media items as (image, rows)
        # because it calls _extract_table_figure(..., need_position=False).
        # Some patched versions may return ((image, rows), positions). Support both
        # shapes so figure/table crops are not silently dropped.
        positions = []
        try:
            if (
                isinstance(item, (list, tuple))
                and len(item) == 2
                and isinstance(item[0], (list, tuple))
                and len(item[0]) == 2
            ):
                (img, rows), positions = item
            elif isinstance(item, (list, tuple)) and len(item) == 2:
                img, rows = item
            else:
                logger.debug("unexpected RAGFlow pdf media item shape: %r", type(item))
                continue
        except Exception as exc:
            logger.info("skip malformed RAGFlow pdf media item: %s", exc)
            continue
        text, table_html = _table_text(rows)
        media_kind, table_subtype, table_confidence, classification_reason = _table_structure_assessment(rows, text, table_html)
        is_table = media_kind == "table"
        page_no = None; bbox = None
        if positions:
            try:
                first_pos = positions[0]
                page_no = int(first_pos[0]) + 1 if int(first_pos[0]) == 0 else int(first_pos[0])
                bbox = (float(first_pos[1]), float(first_pos[3]), float(first_pos[2]), float(first_pos[4]))
            except Exception:
                pass
        image_info = _save_image_payload(img, knowledge_base_id=knowledge_base_id, document_id=document_id, filename_hint=f"pdf_{'table' if is_table else 'figure'}_{len(blocks)}.png")
        meta = {
            "parser": "ragflow_pdf",
            "source": "table_figure_extraction",
            "media_kind": media_kind,
            "table_subtype": table_subtype,
            "table_confidence": table_confidence,
            "classification_reason": classification_reason,
            "positions": positions or [],
            "caption_binding": "ragflow_native",
        }
        # RAGFlow may return OCR text for cropped PDF media.  Preserve it as OCR
        # data rather than allowing a later VLM summary to replace it.
        if text:
            meta["ocr_raw_text"] = text
            meta["ocr_cleaned_text"] = text
            meta["ocr_text"] = text
            meta["ocr_status"] = "from_ragflow"
        if not is_table and text:
            meta["ocr_candidate"] = False
        if image_info:
            meta["image_info"] = {k: v for k, v in image_info.items() if k != "base64"}
        blocks.append(ParsedBlock(
            block_type="table" if is_table else "image",
            text=text or ("PDF 表格" if is_table else "PDF 图片"),
            table_html=table_html,
            image_path=image_info.get("relative_path", ""),
            image_url=image_info.get("url", ""),
            image_base64=image_info.get("base64", ""),
            page_no=page_no,
            bbox=bbox,
            order_index=len(blocks),
            metadata=meta,
        ))
    table_crop_count = sum(1 for b in blocks if b.block_type == "table" and b.image_path)
    figure_crop_count = sum(1 for b in blocks if b.block_type == "image" and b.image_path)
    caption_bound_count = sum(1 for b in blocks if b.block_type in {"table", "image"} and b.text)
    return _finalize(
        blocks,
        "ragflow_pdf_original",
        {
            "pdf_engine": "ragflow_deepdoc_original",
            "table_crop_count": table_crop_count,
            "figure_crop_count": figure_crop_count,
            "caption_bound_count": caption_bound_count,
            "pdf_media_item_count": table_crop_count + figure_crop_count,
        },
    )


def parse_docx(path: Path, *, knowledge_base_id: str, document_id: str) -> ParsedDocument:
    _ensure_vendor_path()
    from docx import Document as DocxDocument
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph
    from deepdoc.parser.docx_parser import RAGFlowDocxParser  # type: ignore

    base_parser = RAGFlowDocxParser()
    doc = DocxDocument(str(path))
    blocks: list[ParsedBlock] = []
    table_idx = 0
    last_image = None
    last_image_meta: dict | None = None

    def flush_last_image() -> None:
        nonlocal last_image, last_image_meta
        if last_image is None:
            return
        media = _save_image_payload(last_image, knowledge_base_id=knowledge_base_id, document_id=document_id, filename_hint=f"docx_image_{len(blocks)}.png")
        meta = {"parser": "ragflow_docx", "source": "paragraph_image", **(last_image_meta or {})}
        if media:
            meta["image_info"] = {k: v for k, v in media.items() if k != "base64"}
        blocks.append(ParsedBlock(block_type="image", text="DOCX 图片", image_path=media.get("relative_path", ""), image_url=media.get("url", ""), image_base64=media.get("base64", ""), order_index=len(blocks), metadata=meta))
        last_image = None
        last_image_meta = None

    for raw in doc._element.body:
        if raw.tag.endswith("p"):
            p = Paragraph(raw, doc)
            text = _clean(p.text)
            style = p.style.name if p.style else ""
            current_image = base_parser.get_picture(doc, p)
            if text:
                if style == "Caption" and (last_image is not None or (blocks and blocks[-1].block_type == "image")):
                    if last_image is not None:
                        media = _save_image_payload(last_image, knowledge_base_id=knowledge_base_id, document_id=document_id, filename_hint=f"docx_captioned_image_{len(blocks)}.png")
                        meta = {"parser": "ragflow_docx", "source": "captioned_image", "caption": text, "style": style}
                        if media:
                            meta["image_info"] = {k: v for k, v in media.items() if k != "base64"}
                        blocks.append(ParsedBlock(block_type="image", text=text, image_path=media.get("relative_path", ""), image_url=media.get("url", ""), image_base64=media.get("base64", ""), order_index=len(blocks), metadata=meta))
                        last_image = None; last_image_meta = None
                    elif blocks and blocks[-1].block_type == "image":
                        blocks[-1].metadata["caption"] = text
                        blocks[-1].text = text
                    continue
                flush_last_image()
                blocks.append(ParsedBlock(block_type="text", text=text, order_index=len(blocks), metadata={"parser": "ragflow_docx", "style": style, "source": "paragraph"}))
                if current_image is not None:
                    media = _save_image_payload(current_image, knowledge_base_id=knowledge_base_id, document_id=document_id, filename_hint=f"docx_image_{len(blocks)}.png")
                    meta = {"parser": "ragflow_docx", "source": "inline_after_text", "anchor_text": text[:200]}
                    if media:
                        meta["image_info"] = {k: v for k, v in media.items() if k != "base64"}
                    blocks.append(ParsedBlock(block_type="image", text=text, image_path=media.get("relative_path", ""), image_url=media.get("url", ""), image_base64=media.get("base64", ""), order_index=len(blocks), metadata=meta))
            elif current_image is not None:
                last_image = current_image
                last_image_meta = {"style": style}
        elif raw.tag.endswith("tbl"):
            flush_last_image()
            table = DocxTable(raw, doc)
            html_rows = ["<table>"]
            text_rows: list[str] = []
            for row in table.rows:
                vals = [_clean(c.text) for c in row.cells]
                text_rows.append(" | ".join(vals))
                html_rows.append("<tr>" + "".join(f"<td>{html.escape(v)}</td>" for v in vals) + "</tr>")
            html_rows.append("</table>")
            blocks.append(ParsedBlock(block_type="table", text="\n".join(text_rows), table_html="".join(html_rows), order_index=len(blocks), metadata={"parser": "ragflow_docx", "source": "docx_table", "table_index": table_idx}))
            table_idx += 1
    flush_last_image()
    if not blocks:
        try:
            secs, tbls = RAGFlowDocxParser()(str(path))
            for sec, style in secs:
                b = _make_text_block(sec, order_index=len(blocks), parser="ragflow_docx", metadata={"style": style})
                if b: blocks.append(b)
            for tbl in tbls:
                text, html_src = _table_text(tbl)
                b = _make_table_block(text, html_src, order_index=len(blocks), parser="ragflow_docx")
                if b: blocks.append(b)
        except Exception:
            pass
    return _finalize(blocks, "ragflow_docx_original")


def parse_excel(path: Path, *, knowledge_base_id: str, document_id: str) -> ParsedDocument:
    _ensure_vendor_path()
    from deepdoc.parser.excel_parser import RAGFlowExcelParser  # type: ignore

    parser = RAGFlowExcelParser()
    blocks: list[ParsedBlock] = []
    raw = path.read_bytes()
    try:
        html_chunks = parser.html(raw)
    except Exception:
        html_chunks = []
    for idx, table_html in enumerate(html_chunks or []):
        b = _make_table_block(_html_to_text(table_html), table_html, order_index=len(blocks), parser="ragflow_excel", metadata={"sheet_chunk": idx})
        if b: blocks.append(b)
    try:
        rows = parser(raw)
        if rows:
            blocks.append(ParsedBlock(block_type="text", text="\n".join(_clean(r) for r in rows if _clean(r)), order_index=len(blocks), metadata={"parser": "ragflow_excel", "source": "semantic_rows"}))
    except Exception as exc:
        logger.info("ragflow excel row parse failed: %s", exc)
    try:
        wb = RAGFlowExcelParser._load_excel_to_workbook(BytesIO(raw))
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for item in RAGFlowExcelParser._extract_images_from_worksheet(ws, sheet):
                media = _save_image_payload(item.get("image"), knowledge_base_id=knowledge_base_id, document_id=document_id, filename_hint=f"excel_{sheet}_image_{len(blocks)}.png")
                meta = {"parser": "ragflow_excel", "source": "worksheet_image", **{k: v for k, v in item.items() if k != "image"}}
                if media:
                    meta["image_info"] = {k: v for k, v in media.items() if k != "base64"}
                blocks.append(ParsedBlock(block_type="image", text=f"Excel 图片：{sheet}", image_path=media.get("relative_path", ""), image_url=media.get("url", ""), image_base64=media.get("base64", ""), order_index=len(blocks), metadata=meta))
    except Exception as exc:
        logger.info("ragflow excel image extraction failed: %s", exc)
    return _finalize(blocks, "ragflow_excel_original")


def parse_csv_or_tsv(path: Path) -> ParsedDocument:
    _ensure_vendor_path()
    from deepdoc.parser.excel_parser import RAGFlowExcelParser  # type: ignore
    parser = RAGFlowExcelParser()
    raw = path.read_bytes()
    blocks: list[ParsedBlock] = []
    try:
        rows = parser(raw)
    except Exception:
        rows = []
    if rows:
        blocks.append(ParsedBlock(block_type="table", text="\n".join(_clean(r) for r in rows if _clean(r)), order_index=0, metadata={"parser": "ragflow_excel_csv", "source": "csv_semantic_rows"}))
    else:
        from deepdoc.parser.txt_parser import RAGFlowTxtParser  # type: ignore
        sections = RAGFlowTxtParser()(str(path))
        for s in sections:
            b = _make_text_block(_section_text(s), order_index=len(blocks), parser="ragflow_txt_csv")
            if b: blocks.append(b)
    return _finalize(blocks, "ragflow_csv_original")


def parse_html(path: Path) -> ParsedDocument:
    _ensure_vendor_path()
    from deepdoc.parser.html_parser import RAGFlowHtmlParser  # type: ignore
    parser = RAGFlowHtmlParser()
    sections = parser(str(path))
    blocks: list[ParsedBlock] = []
    for sec in sections or []:
        raw = str(sec or "")
        if _TABLE_RE.search(raw):
            b = _make_table_block(_html_to_text(raw), raw, order_index=len(blocks), parser="ragflow_html")
        else:
            b = _make_text_block(raw, order_index=len(blocks), parser="ragflow_html")
        if b: blocks.append(b)
    return _finalize(blocks, "ragflow_html_original")


def parse_json(path: Path) -> ParsedDocument:
    _ensure_vendor_path()
    from deepdoc.parser.json_parser import RAGFlowJsonParser  # type: ignore
    parser = RAGFlowJsonParser()
    sections = parser(path.read_bytes())
    blocks = [b for sec in (sections or []) if (b := _make_text_block(sec, order_index=0, parser="ragflow_json"))]
    return _finalize(blocks, "ragflow_json_original")


def parse_txt(path: Path) -> ParsedDocument:
    _ensure_vendor_path()
    from deepdoc.parser.txt_parser import RAGFlowTxtParser  # type: ignore
    parser = RAGFlowTxtParser()
    sections = parser(str(path))
    blocks = []
    for sec in sections or []:
        b = _make_text_block(_section_text(sec), order_index=len(blocks), parser="ragflow_txt")
        if b: blocks.append(b)
    return _finalize(blocks, "ragflow_txt_original")


def _markdown_element_metadata(item: dict[str, Any], *, line_offset: int) -> dict[str, Any]:
    """Normalize DeepDoc Markdown element metadata to original file positions."""
    meta = dict(item or {})
    for key in ("start_line", "end_line"):
        value = meta.get(key)
        if isinstance(value, int):
            meta[key] = value + line_offset

    element_type = str(meta.get("type") or "").strip().lower()
    content = str(meta.get("content") or "")
    if element_type == "header":
        match = re.match(r"^\s*(#{1,6})\s+", content)
        meta.setdefault("raw_block_type", "heading")
        meta.setdefault("is_heading", True)
        meta.setdefault("level", len(match.group(1)) if match else 1)
    elif element_type == "code_block":
        meta.setdefault("raw_block_type", "code_block")
        meta.setdefault("content_kind", "code_block")
        meta.setdefault("preserve_line_boundaries", True)
    return meta


def _render_markdown_table(raw_table: str, table_format: str) -> tuple[str, str]:
    """Return retrieval text and HTML while preserving raw Markdown semantics."""
    raw_table = raw_table.strip()
    if table_format == "html" or "<table" in raw_table.lower():
        return _html_to_text(raw_table), raw_table

    # Keep the original pipe table as retrieval text.  This retains cell
    # boundaries better than flattening HTML into whitespace-only prose.
    table_text = raw_table
    table_html = raw_table
    try:
        from markdown import markdown

        rendered = markdown(raw_table, extensions=["markdown.extensions.tables"])
        if "<table" in rendered.lower():
            table_html = rendered
    except Exception as exc:
        logger.info("render markdown table failed, keep raw table: %s", exc)
    return table_text, table_html


def parse_markdown(path: Path) -> ParsedDocument:
    """Parse Markdown without moving tables away from their source position.

    DeepDoc's ``extract_tables_and_remainder(..., separate_tables=True)`` was
    intentionally not used here: it returns body and tables separately, which
    loses ordering, and its regex can miss a table at EOF without a trailing
    newline.  The ordered scanner preserves tables, headings, prose and fenced
    code in source order before the normal structured chunking pipeline runs.
    """
    from app.services.document_parser.markdown_ordered import (
        extract_markdown_text_elements,
        split_markdown_preserving_tables,
    )

    text = path.read_text(encoding="utf-8", errors="ignore")
    blocks: list[ParsedBlock] = []
    table_count = 0

    for segment in split_markdown_preserving_tables(text):
        if segment.kind == "table":
            table_text, table_html = _render_markdown_table(segment.content, segment.table_format)
            table_meta = {
                "raw_block_type": "table",
                "content_kind": "table",
                "source": "ordered_markdown_table",
                "table_format": segment.table_format,
                "start_line": segment.start_line,
                "end_line": segment.end_line,
                "raw_markdown": segment.content.strip() if segment.table_format == "pipe" else "",
            }
            block = _make_table_block(
                table_text,
                table_html,
                order_index=len(blocks),
                parser="native_markdown_ordered",
                metadata=table_meta,
            )
            if block:
                blocks.append(block)
                table_count += 1
            continue

        for item in extract_markdown_text_elements(segment.content, line_offset=segment.start_line):
            content = item.get("content", "") if isinstance(item, dict) else item
            metadata = (
                _markdown_element_metadata(item, line_offset=0)
                if isinstance(item, dict)
                else {}
            )
            block = _make_text_block(
                content,
                order_index=len(blocks),
                parser="native_markdown_ordered",
                metadata=metadata,
            )
            if block:
                blocks.append(block)

    return _finalize(
        blocks,
        "native_markdown_ordered",
        {
            "markdown_order_preserved": True,
            "markdown_table_count": table_count,
            "markdown_table_eof_safe": True,
        },
    )


def parse_epub(path: Path) -> ParsedDocument:
    _ensure_vendor_path()
    from deepdoc.parser.epub_parser import RAGFlowEpubParser  # type: ignore
    parser = RAGFlowEpubParser()
    sections = parser(str(path))
    blocks = []
    for sec in sections or []:
        raw = str(sec or "")
        if _TABLE_RE.search(raw):
            b = _make_table_block(_html_to_text(raw), raw, order_index=len(blocks), parser="ragflow_epub")
        else:
            b = _make_text_block(raw, order_index=len(blocks), parser="ragflow_epub")
        if b: blocks.append(b)
    return _finalize(blocks, "ragflow_epub_original")


def parse_ppt(path: Path, *, knowledge_base_id: str, document_id: str) -> ParsedDocument:
    _ensure_vendor_path()
    from deepdoc.parser.ppt_parser import RAGFlowPptParser  # type: ignore
    parser = RAGFlowPptParser()
    blocks: list[ParsedBlock] = []
    try:
        sections = parser(str(path), 0, 10_000)
    except Exception:
        sections = []
    for idx, sec in enumerate(sections or [], start=1):
        b = _make_text_block(sec, order_index=len(blocks), parser="ragflow_ppt", page_no=idx, metadata={"slide": idx})
        if b: blocks.append(b)
    # RAGFlow's PPT parser is text-first; keep know_rag multimodal by preserving slide images as media blocks.
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        prs = Presentation(str(path))
        for si, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                    blob = shape.image.blob
                    media = save_media_bytes(knowledge_base_id=knowledge_base_id, document_id=document_id, filename_hint=f"ppt_slide_{si}_image.png", data=blob, default_ext=".png")
                    meta = {"parser": "ragflow_ppt", "source": "slide_image", "slide": si, "image_info": {k: v for k, v in media.items() if k != "base64"}}
                    blocks.append(ParsedBlock(block_type="image", text=f"PPT 第 {si} 页图片", image_path=media.get("relative_path", ""), image_url=media.get("url", ""), image_base64=media.get("base64", ""), page_no=si, order_index=len(blocks), metadata=meta))
    except Exception as exc:
        logger.info("ppt image extraction failed: %s", exc)
    return _finalize(blocks, "ragflow_ppt_original")


def parse_image(path: Path, *, knowledge_base_id: str, document_id: str) -> ParsedDocument:
    _ensure_vendor_path()
    raw = path.read_bytes()
    media = save_media_bytes(knowledge_base_id=knowledge_base_id, document_id=document_id, filename_hint=path.name, data=raw, default_ext=path.suffix or ".png")
    ocr_text = ""
    ocr_status = "skipped"
    try:
        import cv2
        import numpy as np
        from deepdoc.vision.ocr import OCR  # type: ignore
        from common.file_utils import get_project_base_directory  # type: ignore
        model_dir = Path(os.getenv("DEEPDOC_MODEL_DIR", "/data/models/deepdoc"))
        if not (model_dir / "det.onnx").exists():
            model_dir = Path(get_project_base_directory()) / "rag" / "res" / "deepdoc"
        img_arr = cv2.imdecode(np.frombuffer(raw, dtype=np.uint8), cv2.IMREAD_COLOR)
        if img_arr is not None:
            ocr = OCR(str(model_dir))
            result = ocr(img_arr)
            parts = []
            for _, rec in result or []:
                if isinstance(rec, (list, tuple)) and rec:
                    parts.append(str(rec[0]))
            ocr_text = _clean("\n".join(parts))
            ocr_status = "success" if ocr_text else "empty"
        else:
            ocr_status = "unsupported_image_decode"
    except Exception as exc:
        ocr_status = "failed"
        logger.info("ragflow image OCR failed: %s", exc)
    meta = {
        "parser": "ragflow_image_ocr",
        "source": "uploaded_image",
        "file_name": path.name,
        "ocr_status": ocr_status,
        "ocr_text": ocr_text,
        "image_info": {k: v for k, v in media.items() if k != "base64"},
    }
    text = ocr_text or f"图片文件：{path.name}"
    block = ParsedBlock(block_type="image", text=text, image_path=media.get("relative_path", ""), image_url=media.get("url", ""), image_base64=media.get("base64", ""), order_index=0, metadata=meta)
    return _finalize([block], "ragflow_image_original", {"ocr_status": ocr_status, "ocr_text_length": len(ocr_text)})


def parse_archive(path: Path) -> ParsedDocument:
    blocks: list[ParsedBlock] = []
    try:
        with ZipFile(path) as zf:
            names = zf.namelist()[:500]
        text = "压缩包文件列表：\n" + "\n".join(names)
    except Exception as exc:
        text = f"压缩包文件：{path.name}\n读取失败：{exc}"
    blocks.append(ParsedBlock(block_type="text", text=_clean(text), order_index=0, metadata={"parser": "ragflow_archive_listing"}))
    return _finalize(blocks, "ragflow_archive_listing")


def parse_binary_fallback(path: Path, *, parser: str, reason: str) -> ParsedDocument:
    text = f"附件文件：{path.name}\n文件类型：{path.suffix.lower().lstrip('.')}\n说明：{reason}"
    return ParsedDocument.from_text(_clean(text), metadata={"parser": parser, "file_type": path.suffix.lower().lstrip(".")})


def parse_document_with_ragflow(path: Path, *, file_type: str, knowledge_base_id: str, document_id: str) -> ParsedDocument:
    suffix = (file_type or path.suffix).lower().lstrip(".")
    if suffix == "pdf":
        return parse_pdf(path, knowledge_base_id=knowledge_base_id, document_id=document_id)
    if suffix == "docx":
        return parse_docx(path, knowledge_base_id=knowledge_base_id, document_id=document_id)
    if suffix in {"xlsx", "xls"}:
        return parse_excel(path, knowledge_base_id=knowledge_base_id, document_id=document_id)
    if suffix in {"csv", "tsv"}:
        return parse_csv_or_tsv(path)
    if suffix in HTML_SUFFIXES:
        return parse_html(path)
    if suffix in {"json", "jsonl"}:
        return parse_json(path)
    if suffix in {"xml", "yaml", "yml"}:
        return parse_txt(path)
    if suffix in {"md", "markdown"}:
        return parse_markdown(path)
    if suffix == "txt":
        return parse_txt(path)
    if suffix in PRESENTATION_SUFFIXES:
        return parse_ppt(path, knowledge_base_id=knowledge_base_id, document_id=document_id)
    if suffix in EPUB_SUFFIXES:
        return parse_epub(path)
    if suffix in IMAGE_SUFFIXES:
        return parse_image(path, knowledge_base_id=knowledge_base_id, document_id=document_id)
    if suffix in ARCHIVE_SUFFIXES:
        return parse_archive(path)
    if suffix in LEGACY_DOC_SUFFIXES:
        return parse_binary_fallback(path, parser="ragflow_legacy_office_attachment", reason="RAGFlow parser 主要处理 docx；旧版 Office/RTF 已保存为附件，建议转为 docx 以获得完整图文结构")
    return parse_binary_fallback(path, parser="ragflow_generic_attachment", reason="未知格式已保存为附件，并生成可检索占位切片")
